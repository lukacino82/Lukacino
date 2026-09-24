"""Open ± X fade: limitní long na open - X, limitní short na open + X.
TP / SL v tickách, výstup na TP, SL nebo na konci seance.

Použití:
  python open_fade/open_fade.py DATA.txt --levels 10 15 20 40 --tp 60 --sl 40
DATA = export ze Sierra Chart (Date, Time, Open, High, Low, Last, ...) v čase ET,
       buď .txt/.csv, nebo .xlsx (bere se první list s hlavičkou Date/Time).
       Funguje pro časové i volume bary (čas = začátek baru).
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass

import numpy as np
import pandas as pd

TICK = 0.25
POINT_USD = 50.0  # ES; MES = 5.0


@dataclass
class Params:
    level: float            # X v bodech od open
    tp_ticks: int = 60
    sl_ticks: int = 40
    session_open: str = "09:30"
    last_entry: str = "15:59"   # do kdy smí limitka vstoupit
    session_close: str = "16:00"
    cost_pts: float = 0.0   # náklady na round-trip v bodech


def load_bars(path: str) -> pd.DataFrame:
    if path.lower().endswith((".xlsx", ".xlsm")):
        import datetime as dt
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = next(w for w in wb.worksheets
                  if [str(c).strip().lower() for c in next(w.iter_rows(max_row=1, values_only=True))[:2]] == ["date", "time"])
        rows = [r[:6] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        df = pd.DataFrame(rows, columns=["date", "time", "open", "high", "low", "close"])
        df.index = pd.DatetimeIndex([dt.datetime.combine(d.date(), t) for d, t in zip(df.date, df.time)])
        df = df.drop(columns=["date", "time"])
    elif path.lower().endswith(".pkl"):
        df = pd.read_pickle(path)[["open", "high", "low", "close"]]
    else:
        raw = pd.read_csv(path, skipinitialspace=True)
        raw.columns = [c.strip().lower() for c in raw.columns]
        ts = pd.to_datetime(raw["date"].astype(str).str.strip() + " " + raw["time"].astype(str).str.strip(),
                            format="mixed")
        df = pd.DataFrame({"open": raw.open, "high": raw.high, "low": raw.low, "close": raw["last"]})
        df.index = pd.DatetimeIndex(ts)
    # volume bary mohou mít stejný čas – pořadí zachovat, nic nemazat
    return df.astype(float).sort_index(kind="stable")


def simulate_day(o, h, l, c, tm, p: Params, direction: int, optimistic: bool = False):
    """Vrátí (výsledek, body) pro jeden směr jednoho dne, nebo None, pokud limitka nevstoupila.
    direction +1 = long na open - X, -1 = short na open + X."""
    tp, sl = p.tp_ticks * TICK, p.sl_ticks * TICK
    day_open = o[0]
    level = day_open - direction * p.level
    last_entry = pd.Timestamp(p.last_entry).time()
    for i in range(len(o)):
        if tm[i] > last_entry:
            return None
        touched = l[i] <= level if direction > 0 else h[i] >= level
        if not touched:
            continue
        # gap přes úroveň -> plnění na open baru (lepší cena)
        entry = min(level, o[i]) if direction > 0 else max(level, o[i])
        stop, target = entry - direction * sl, entry + direction * tp
        for j in range(i, len(o)):
            hit_sl = l[j] <= stop if direction > 0 else h[j] >= stop
            hit_tp = (h[j] >= target if direction > 0 else l[j] <= target) and j > i
            if j > i:  # gap přes SL/TP na open baru
                if (o[j] <= stop if direction > 0 else o[j] >= stop):
                    return "SL", direction * (o[j] - entry)
                if (o[j] >= target if direction > 0 else o[j] <= target):
                    return "TP", direction * (o[j] - entry)
            if hit_sl and hit_tp and optimistic:
                return "TP", tp
            if hit_sl:  # SL i TP na stejném baru -> konzervativně SL
                return "SL", -sl
            if hit_tp:
                return "TP", tp
        return "EOD", direction * (c[-1] - entry)
    return None


def run(df: pd.DataFrame, p: Params, optimistic: bool = False) -> pd.DataFrame:
    t = df.index.time
    rth = df[(t >= pd.Timestamp(p.session_open).time()) & (t < pd.Timestamp(p.session_close).time())]
    rows = []
    for d, g in rth.groupby(rth.index.date):
        if g.index[-1].time() < pd.Timestamp("15:45").time():  # zkrácené seance (svátky)
            continue
        arr = (g.open.values, g.high.values, g.low.values, g.close.values, g.index.time)
        for direction, name in ((1, "LONG open-X"), (-1, "SHORT open+X")):
            r = simulate_day(*arr, p, direction, optimistic)
            if r:
                rows.append(dict(date=pd.Timestamp(d), side=name, result=r[0], pts=r[1] - p.cost_pts))
    return pd.DataFrame(rows)


def summarize(tr: pd.DataFrame, n_days: int, p: Params) -> pd.DataFrame:
    out = []
    for side, g in list(tr.groupby("side")) + [("OBA SMĚRY", tr)]:
        n = len(g); tp = (g.result == "TP").sum(); sl = (g.result == "SL").sum(); eod = (g.result == "EOD").sum()
        pts = g.pts
        out.append({
            "X": p.level, "směr": side, "dní": n_days, "vstupů": n,
            "vstup % dní": 100 * n / n_days / (2 if side == "OBA SMĚRY" else 1),
            "TP": tp, "SL": sl, "konec seance": eod,
            "TP % (z TP+SL)": 100 * tp / max(tp + sl, 1),
            "TP % (ze všech)": 100 * tp / max(n, 1),
            "Ø konec seance b.": g[g.result == "EOD"].pts.mean() if eod else np.nan,
            "b./obchod": pts.mean(), "t": pts.mean() / pts.std() * np.sqrt(n) if n > 1 else np.nan,
            "$ celkem (1 ES)": pts.sum() * POINT_USD,
            "PF": pts[pts > 0].sum() / -pts[pts < 0].sum() if (pts < 0).any() else np.inf,
        })
    return pd.DataFrame(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--levels", type=float, nargs="+", default=[10, 15, 20, 40])
    ap.add_argument("--tp", type=int, default=60, help="TP v tickách")
    ap.add_argument("--sl", type=int, default=40, help="SL v tickách")
    ap.add_argument("--cost", type=float, default=0.0, help="náklady round-trip v bodech")
    ap.add_argument("--last-entry", default="15:59")
    ap.add_argument("--optimistic", action="store_true", help="SL i TP na stejném baru počítat jako TP")
    ap.add_argument("--by-year", action="store_true", help="rozpad po letech")
    ap.add_argument("--out", default=None, help="CSV se souhrnem")
    a = ap.parse_args()
    df = load_bars(a.path)
    t = df.index.time
    rth = df[(t >= pd.Timestamp("09:30").time()) & (t < pd.Timestamp("16:00").time())]
    last = rth.groupby(rth.index.date).apply(lambda g: g.index[-1].time())
    n_days = int((last >= pd.Timestamp("15:45").time()).sum())
    res, trades = [], []
    for x in a.levels:
        p = Params(level=x, tp_ticks=a.tp, sl_ticks=a.sl, cost_pts=a.cost, last_entry=a.last_entry)
        tr = run(df, p, a.optimistic); tr["X"] = x; trades.append(tr)
        res.append(summarize(tr, n_days, p))
    R = pd.concat(res, ignore_index=True)
    pd.set_option("display.width", 250)
    print(f"RTH dní: {n_days}, TP {a.tp} ticků = {a.tp*TICK} b., SL {a.sl} ticků = {a.sl*TICK} b., "
          f"RRR {a.tp/a.sl:.2f}, break-even TP% = {100*a.sl/(a.tp+a.sl):.1f} %, náklady {a.cost} b.")
    print(f"Období: {df.index[0].date()} až {df.index[-1].date()}" + ("  [OPTIMISTICKÁ varianta]" if a.optimistic else ""))
    print(R.round(2).to_string(index=False))
    if a.by_year:
        T = pd.concat(trades)
        y = T.groupby(["X", "side", T.date.dt.year]).pts.sum().unstack().round(0)
        print("\nBody po letech (součet):"); print(y.to_string())
    if a.out:
        R.to_csv(a.out, index=False)
        pd.concat(trades).to_csv(a.out.replace(".csv", "_obchody.csv"), index=False)


if __name__ == "__main__":
    main()
